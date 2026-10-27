#!/usr/bin/env python3
"""Find and download legally open full texts for the 189-report review queue.

Run from a computer with internet access:
  python -m pip install requests openpyxl pypdf
  set PAPER_REVIEW_UNPAYWALL_EMAIL=you@example.org   # Windows PowerShell: $env:PAPER_REVIEW_UNPAYWALL_EMAIL='...'
  python fetch_189_open_access.py

The script uses public metadata/OA APIs only. It does not bypass paywalls.
It writes PDFs to 189条全文PDF/ and updates 189条全文批量获取清单.xlsx.
"""
from __future__ import annotations
import os, re, sys, time, json, hashlib, urllib.request, urllib.error, urllib.parse
from pathlib import Path
from datetime import datetime, timezone
from urllib.parse import quote, quote_plus

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print('Missing dependency:', e)
    print('Install with: python -m pip install openpyxl pypdf')
    raise SystemExit(2)

BASE = Path(__file__).resolve().parent
INPUT = BASE / '189条全文批量获取清单.xlsx'
PDF_DIR = BASE / '189条全文PDF'
EMAIL = os.getenv('PAPER_REVIEW_UNPAYWALL_EMAIL', '').strip()
TIMEOUT = 25
DISABLED_APIS = set()

class HTTPResponse:
    def __init__(self, data, url, headers, status):
        self.data=data; self.url=url; self.headers=dict(headers.items()); self.status_code=status
    def raise_for_status(self):
        if self.status_code>=400: raise RuntimeError(f'HTTP {self.status_code}: {self.url}')
    def json(self): return json.loads(self.data.decode('utf-8'))
    def iter_content(self, size):
        for i in range(0,len(self.data),size): yield self.data[i:i+size]

class HTTPSession:
    headers={'User-Agent':'LLMExercisePrescriptionReview/1.0 (lawful OA full text retrieval)','Accept':'application/json,application/pdf,*/*'}
    def get(self,url,params=None,timeout=TIMEOUT,stream=False,allow_redirects=True):
        if params:
            join='&' if '?' in url else '?'
            url += join + urllib.parse.urlencode(params)
        req=urllib.request.Request(url,headers=self.headers)
        try:
            with urllib.request.urlopen(req,timeout=timeout) as response:
                return HTTPResponse(response.read(),response.geturl(),response.headers,response.status)
        except urllib.error.HTTPError as e:
            raise RuntimeError(f'HTTP {e.code}: {url}')
        except Exception as e:
            raise RuntimeError(f'{type(e).__name__}: {e}')

SESSION=HTTPSession()

def norm(s): return re.sub(r'[^a-z0-9]+',' ',(s or '').lower()).strip()
def doi_clean(s): return re.sub(r'^https?://(dx\.)?doi\.org/','',(s or '').strip(),flags=re.I).rstrip(' .')
def safe_name(s): return re.sub(r'[^A-Za-z0-9_-]+','_',s or 'record')
def http_json(url, params=None):
    service='OpenAlex' if 'api.openalex.org' in url else ('Europe PMC' if 'europepmc' in url else ('Unpaywall' if 'unpaywall.org' in url else 'other'))
    if service in DISABLED_APIS: raise RuntimeError(service+' disabled after a prior 403/proxy/network failure')
    try:r=SESSION.get(url,params=params,timeout=TIMEOUT)
    except Exception as e:
        msg=str(e)
        if any(code in msg for code in ('401','403','407','URLError','Tunnel')): DISABLED_APIS.add(service)
        raise RuntimeError(f'{service} stopped after interface/network failure: {msg}')
    if r.status_code in (401,403,407):
        DISABLED_APIS.add(service)
        raise RuntimeError(f'{service} stopped after HTTP {r.status_code}; no further requests will be sent to this interface in this run')
    try:r.raise_for_status()
    except Exception:
        DISABLED_APIS.add(service)
        raise
    return r.json(),r.url
def openalex_find(title,doi):
    if doi:
        try:
            data,url=http_json('https://api.openalex.org/works/https://doi.org/'+quote(doi,safe='/'))
            if isinstance(data,dict) and data.get('id'): return data,url
        except Exception: pass
    data,url=http_json('https://api.openalex.org/works',{'search':title,'per-page':5})
    items=data.get('results',[])
    target=norm(title)
    for item in items:
        cand=norm(item.get('title',''))
        if cand and (cand==target or (len(target)>35 and (target in cand or cand in target))):
            return item,url
    return None,url
def europepmc_find(title,doi,pmid):
    if pmid and str(pmid).isdigit():
        data,url=http_json('https://www.ebi.ac.uk/europepmc/webservices/rest/search',{'query':'EXT_ID:'+str(pmid),'format':'json','resultType':'core'})
        hits=data.get('resultList',{}).get('result',[])
        if hits:return hits[0],url
    if doi:
        data,url=http_json('https://www.ebi.ac.uk/europepmc/webservices/rest/search',{'query':'EXT_ID:'+doi,'format':'json','resultType':'core'})
        hits=data.get('resultList',{}).get('result',[])
        if hits:return hits[0],url
    data,url=http_json('https://www.ebi.ac.uk/europepmc/webservices/rest/search',{'query':'TITLE:"'+title.replace('"','')+'"','format':'json','resultType':'core','pageSize':5})
    target=norm(title)
    for hit in data.get('resultList',{}).get('result',[]):
        if norm(hit.get('title',''))==target:return hit,url
    return None,url
def extract_openalex_pdf(item):
    if not item:return []
    urls=[]
    for loc in item.get('locations',[]) or []:
        src=loc.get('source') or {}
        if loc.get('is_oa') or (src.get('is_oa') if src else False):
            u=loc.get('pdf_url') or (loc.get('landing_page_url') if str(loc.get('landing_page_url','')).lower().endswith('.pdf') else None)
            if u:urls.append((u,'OpenAlex OA location'))
    best=item.get('best_oa_location') or {}
    if best.get('pdf_url'):urls.insert(0,(best['pdf_url'],'OpenAlex best OA PDF'))
    return list(dict.fromkeys(urls))
def extract_epmc_pdf(hit):
    if not hit:return []
    arr=[]
    for x in (hit.get('fullTextUrlList') or {}).get('fullTextUrl',[]) or []:
        u=x.get('url','')
        if x.get('documentStyle')=='pdf' or u.lower().endswith('.pdf'):
            arr.append((u,'Europe PMC fullTextUrl PDF'))
    pmcid=hit.get('pmcid')
    if pmcid: arr.insert(0,(f'https://europepmc.org/articles/{pmcid}?pdf=render','Europe PMC PMC PDF'))
    return list(dict.fromkeys(arr))
def unpaywall_pdf(doi):
    if not doi or not EMAIL:return None
    data,url=http_json('https://api.unpaywall.org/v2/'+quote(doi,safe=''),{'email':EMAIL})
    loc=data.get('best_oa_location') or {}
    u=loc.get('url_for_pdf')
    if u:return (u,'Unpaywall best OA PDF'),url
    return None,url
def verify_pdf(path,title,doi):
    text=''
    try:
        from pypdf import PdfReader
        reader=PdfReader(str(path)); text=' '.join((p.extract_text() or '') for p in reader.pages[:3])
    except Exception:
        try:
            import subprocess
            text=subprocess.run(['pdftotext','-f','1','-l','3',str(path),'-'],capture_output=True,text=True,timeout=20).stdout
        except Exception: pass
    nt=norm(text); ntitle=norm(title)
    title_hit=bool(ntitle and (ntitle[:60] in nt or nt[:80] in ntitle))
    doi_hit=bool(doi and doi.lower() in text.lower())
    return (title_hit or doi_hit),title_hit,doi_hit
def download(url,record_id,title,doi):
    r=SESSION.get(url,timeout=TIMEOUT,stream=True,allow_redirects=True)
    r.raise_for_status()
    ctype=r.headers.get('content-type','').lower()
    chunks=[]; total=0
    for chunk in r.iter_content(1024*128):
        if chunk:
            total+=len(chunk)
            if total>40*1024*1024: raise ValueError('PDF exceeds 40 MB safety limit')
            chunks.append(chunk)
    blob=b''.join(chunks)
    if not blob.startswith(b'%PDF-') and 'pdf' not in ctype: raise ValueError('Response is not a PDF')
    tmp=PDF_DIR/(safe_name(record_id)+'.download')
    tmp.write_bytes(blob)
    ok,th,dh=verify_pdf(tmp,title,doi)
    if not ok:
        tmp.unlink(missing_ok=True)
        raise ValueError('PDF identity check failed (title/DOI absent in first three pages)')
    dest=PDF_DIR/(safe_name(record_id)+'.pdf');tmp.replace(dest)
    return dest,hashlib.sha256(blob).hexdigest(),th,dh,r.url

def main():
    if not INPUT.exists(): raise SystemExit(f'Input workbook not found: {INPUT}')
    PDF_DIR.mkdir(exist_ok=True)
    wb=load_workbook(INPUT)
    ws=wb['189条获取清单']
    headers={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
    statuscol=headers['获取状态']; n=0
    for row in range(2,ws.max_row+1):
        rid=ws.cell(row,headers['Record_ID']).value
        if not rid: continue
        title=str(ws.cell(row,headers['题名']).value or '').strip()
        doi=doi_clean(str(ws.cell(row,headers['DOI']).value or ''))
        pmid=str(ws.cell(row,headers['PMID']).value or '').strip()
        sources=[]; pdf_candidates=[]; errors=[]; oa_item=None; epmc_hit=None
        ws.cell(row,headers['执行日期'],datetime.now(timezone.utc).astimezone().strftime('%Y-%m-%d %H:%M %Z'))
        try:
            oa_item,oa_url=openalex_find(title,doi); sources.append('OpenAlex: '+oa_url)
            pdf_candidates += extract_openalex_pdf(oa_item)
        except Exception as e: errors.append('OpenAlex: '+str(e))
        try:
            epmc_hit,epmc_url=europepmc_find(title,doi,pmid); sources.append('Europe PMC: '+epmc_url)
            pdf_candidates += extract_epmc_pdf(epmc_hit)
        except Exception as e: errors.append('Europe PMC: '+str(e))
        if doi and EMAIL:
            try:
                got,uwurl=unpaywall_pdf(doi);sources.append('Unpaywall: '+uwurl)
                if got:pdf_candidates.append(got)
            except Exception as e:errors.append('Unpaywall: '+str(e))
        # DOI identity is checked exactly. Title-only matches require author/year corroboration.
        api_titles=[]
        if oa_item and oa_item.get('title'):api_titles.append(oa_item['title'])
        if epmc_hit and epmc_hit.get('title'):api_titles.append(epmc_hit['title'])
        author_year=[]
        if oa_item:
            auth=[(a.get('author') or {}).get('display_name','') for a in oa_item.get('authorships',[]) or []]
            if auth: author_year.append('OpenAlex: '+', '.join(x for x in auth[:3] if x))
        if epmc_hit:
            au=epmc_hit.get('authorString')
            if au:author_year.append('Europe PMC: '+au)
        if author_year:ws.cell(row,headers['在线元数据作者']).value='; '.join(author_year)
        years=[]
        if oa_item and oa_item.get('publication_year'):years.append(str(oa_item['publication_year']))
        if epmc_hit and epmc_hit.get('firstPublicationDate'):years.append(str(epmc_hit['firstPublicationDate'])[:4])
        if years:ws.cell(row,headers['在线元数据年份']).value='; '.join(dict.fromkeys(years))
        api_dois=[]
        if oa_item and oa_item.get('doi'):api_dois.append(doi_clean(oa_item.get('doi')))
        if epmc_hit and epmc_hit.get('doi'):api_dois.append(doi_clean(epmc_hit.get('doi')))
        doi_identity=bool(doi and any(x.lower()==doi.lower() for x in api_dois))
        title_identity=any(norm(t)==norm(title) for t in api_titles)
        input_author=str(ws.cell(row,headers['作者']).value or '').strip()
        input_year=str(ws.cell(row,headers['年份']).value or '').strip()
        author_text=' '.join(author_year).lower()
        author_identity=bool(input_author and input_author not in ('NR','NR—源记录未提供') and input_author.lower() in author_text)
        year_identity=bool(input_year and any(input_year==y for y in years))
        identity=doi_identity or (title_identity and author_identity and year_identity)
        ws.cell(row,headers['身份匹配状态'],'MATCHED_DOI' if doi_identity else ('MATCHED_TITLE_AUTHOR_YEAR' if identity else ('METADATA_FOUND_NEEDS_REVIEW' if api_titles else 'NO_METADATA_MATCH')))
        if errors and not api_titles:
            status='网络访问失败'
        elif not identity:
            status='元数据待核实'
        elif pdf_candidates:
            status='找到全文链接但未下载'
            for url,src in dict.fromkeys(pdf_candidates):
                try:
                    dest,sha,th,dh,finalurl=download(url,rid,title,doi)
                    status='PDF已下载';ws.cell(row,headers['本地PDF路径'],str(dest));ws.cell(row,headers['SHA256'],sha)
                    ws.cell(row,headers['PDF题名核验'],'MATCH' if th else 'DOI_MATCH')
                    ws.cell(row,headers['PDF来源'],src+' | '+finalurl)
                    break
                except Exception as e:errors.append(src+': '+str(e))
        else:
            # Distinguish repository metadata with no free PDF from no reliable result.
            has_record=bool(oa_item or epmc_hit)
            status='需机构访问' if has_record else ('网络访问失败' if errors else '查无可靠来源')
        ws.cell(row,statuscol,status)
        ws.cell(row,headers['实际来源'], ' | '.join(sources+errors))
        ws.cell(row,headers['检索日期'],datetime.now(timezone.utc).strftime('%Y-%m-%d'))
        n+=1
        if n%10==0:
            wb.save(INPUT)
            print(f'Processed {n}/{ws.max_row-1}')
        time.sleep(.15)
    wb.save(INPUT)
    print(f'Done: {n} rows. Workbook: {INPUT}; PDFs: {PDF_DIR}')

if __name__=='__main__': main()
